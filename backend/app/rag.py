"""Ingestion + hybrid retrieval (same interface as the original Chroma version).

Chunks and their embedding vectors live in the platform database (SQLite
locally, PostgreSQL in production) so documents and vectors survive restarts,
redeploys and logins. Retrieval keeps the original pipeline:

    dense (vector cosine) + sparse (BM25)  →  Reciprocal Rank Fusion  →  rerank

Every query is restricted to the doc_ids the caller is allowed to see
(per-document / per-department ACL enforced upstream in main.py).
"""
import io, csv, json, math, re
from pathlib import Path
from . import db as D, embeddings as E
from .config import MIN_RERANK  # noqa: F401  (kept for parity with the original module)

_tok = lambda s: re.findall(r"\w+", s.lower())

# Question stopwords for the sufficiency gate only (BM25 retrieval keeps the
# raw tokens — this list exists so function words don't dilute coverage).
_STOP = {
    "the", "and", "for", "with", "from", "that", "this", "these", "those", "than", "then",
    "what", "when", "where", "which", "who", "whom", "whose", "why", "how",
    "does", "did", "doing", "is", "are", "was", "were", "be", "been", "being",
    "can", "could", "should", "would", "will", "shall", "may", "might", "must",
    "have", "has", "had", "about", "into", "onto", "them", "they", "their", "there",
    "you", "your", "our", "its", "not", "any", "all", "also", "much", "many", "most",
    "some", "such", "per", "via", "over", "under", "between", "during", "before", "after",
}


def _content_tokens(text: str) -> set[str]:
    return {t for t in _tok(text) if len(t) >= 3 and t not in _STOP}


def sufficiency(query: str, hits: list[dict]) -> dict:
    """Deterministic retrieval-quality signal (no LLM, no network).

    Returns the fraction of the question's content terms present in the
    retrieved evidence pool. The agent uses this to decide whether evidence
    is strong enough to attempt generation (else: one controlled retry, then
    abstention). Purely lexical by design — measurable, cheap, testable.
    """
    q = _content_tokens(query)
    if not q:
        return {"coverage": 1.0 if hits else 0.0, "hit_terms": 0, "query_terms": 0}
    if not hits:
        return {"coverage": 0.0, "hit_terms": 0, "query_terms": len(q)}
    ev: set[str] = set()
    for h in hits:
        ev |= _content_tokens(h["text"])
    n = len(q & ev)
    return {"coverage": round(n / len(q), 4), "hit_terms": n, "query_terms": len(q)}


def extract(filename: str, data: bytes) -> list[tuple[int, str]]:
    """Return [(page_number, text)] for the uploaded file."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)]
    if ext == ".docx":
        # Paragraphs AND table cells, in document order: table-hosted facts
        # (policy limits, contact lists, SLAs) were previously invisible.
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(io.BytesIO(data))
        parts: list[str] = []
        for child in d.element.body.iterchildren():
            if child.tag.endswith("}p"):
                parts.append(Paragraph(child, d).text)
            elif child.tag.endswith("}tbl"):
                for row in Table(child, d).rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        return [(1, "\n".join(x for x in parts if x.strip()))]
    if ext == ".pptx":
        # Slide text AND table shapes (sh.has_table), per slide.
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides):
            parts = []
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        parts.append(" | ".join(c.text_frame.text for c in row.cells))
            out.append((i + 1, "\n".join(x for x in parts if x.strip())))
        return out
    text = data.decode("utf-8", errors="ignore")
    if ext == ".csv":
        text = "\n".join(", ".join(r) for r in csv.reader(io.StringIO(text)))
    if ext == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
    return [(1, text)]


def chunk(text, size=900, overlap=150):
    paras, out, cur = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()], [], ""
    for p in paras:
        while len(p) > size:  # hard-split very long paragraphs
            out.append((cur + " " + p[:size]).strip()); cur = ""; p = p[size - overlap:]
        if len(cur) + len(p) > size and cur:
            out.append(cur.strip()); cur = cur[-overlap:]
        cur += " " + p
    if cur.strip(): out.append(cur.strip())
    return out


def ingest(tid, doc_id, filename, data) -> int:
    """Extract → chunk → embed → store. Returns the number of chunks."""
    rows = []
    for page, text in extract(filename, data):
        for j, ch in enumerate(chunk(text)):
            rows.append((f"{doc_id}:{page}:{j}", tid, doc_id, page, ch))
    if not rows:
        return 0
    vectors = E.embed([r[4] for r in rows])          # None → BM25-only mode
    if vectors is not None and len(vectors) != len(rows):
        vectors = None                               # never store a partial/misaligned batch
    emb = vectors if vectors is not None else [None] * len(rows)
    with D.db() as c:
        c.executemany(
            "INSERT INTO chunks(id,tenant_id,doc_id,page,text,embedding) VALUES(?,?,?,?,?,?)",
            [r + (None if v is None else json.dumps(v),) for r, v in zip(rows, emb)],
        )
    return len(rows)


def delete_doc(tid, doc_id):
    with D.db() as c:
        c.execute("DELETE FROM chunks WHERE tenant_id=? AND doc_id=?", (tid, doc_id))


def backfill_embeddings(tid=None, doc_ids=None, batch=32, limit=0):
    """Embed any stored chunks that are missing vectors (idempotent).

    Documents ingested while no embedding key was configured have
    ``embedding = NULL`` and stay BM25-only forever — adding a key later does
    NOT retroactively vectorize them. This walks the chunks table (optionally
    scoped to a tenant/document set), embeds the missing texts with the
    CURRENT provider/model, and writes the vectors back. ``limit`` caps how
    many chunks are processed in this call (0 = all) so callers can page
    through large backlogs inside serverless timeouts. Returns a stats dict;
    never raises on provider errors (reports ``embedded: 0`` and stops).

    Scope note: this is a data-admin operation. Callers must only pass
    tenants/doc_ids they administer; it performs no authorization of its own.
    """
    where, params = ["embedding IS NULL"], []
    if tid is not None:
        where.append("tenant_id=?"); params.append(tid)
    if doc_ids is not None:
        if not doc_ids:
            return {"pending": 0, "embedded": 0, "provider": E.provider()}
        where.append(f"doc_id IN ({','.join('?' * len(doc_ids))})"); params += list(doc_ids)
    sql = f"SELECT id, tenant_id, text FROM chunks WHERE {' AND '.join(where)} ORDER BY id"
    if limit:
        sql += f" LIMIT {int(limit)}"
    with D.db() as c:
        rows = c.execute(sql, params).fetchall()
    pending = len(rows)
    if not pending:
        return {"pending": 0, "embedded": 0, "provider": E.provider()}
    done = 0
    for i in range(0, pending, batch):
        part = rows[i:i + batch]
        vecs = E.embed([r["text"] for r in part])
        if vecs is None or len(vecs) != len(part):
            break                                    # provider gone/failed: stop, keep the rest pending
        with D.db() as c:
            c.executemany("UPDATE chunks SET embedding=? WHERE id=?",
                          [(json.dumps(v), r["id"]) for r, v in zip(part, vecs)])
        done += len(part)
    return {"pending": pending, "embedded": done, "provider": E.provider()}


def drop_tenant(tid):
    with D.db() as c:
        c.execute("DELETE FROM chunks WHERE tenant_id=?", (tid,))


def _cosine(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a)); nb = math.sqrt(sum(y * y for y in b))
    return dot / (na * nb) if na and nb else 0.0


def search(tid, doc_ids, query, k=8):
    """Hybrid: dense + BM25 fused with Reciprocal Rank Fusion, scoped to doc_ids."""
    if not doc_ids:
        return []
    # Chunk the IN (...) list to stay well under statement parameter limits.
    rows = []
    for i in range(0, len(doc_ids), 400):
        part = doc_ids[i:i + 400]
        with D.db() as c:
            rows += c.execute(
                f"SELECT ch.id, ch.text, ch.page, ch.doc_id, ch.embedding, d.filename "
                f"FROM chunks ch JOIN documents d ON d.id=ch.doc_id "
                f"WHERE ch.tenant_id=? AND ch.doc_id IN ({','.join('?' * len(part))})",
                (tid, *part),
            ).fetchall()
    if not rows:
        return []

    # id -> (text, meta) in the same shape the agent expects
    info = {
        r["id"]: (r["text"], {"doc_id": r["doc_id"], "filename": r["filename"], "page": r["page"]})
        for r in rows
    }
    order = [r["id"] for r in rows]
    vectors = {}
    for r in rows:
        if r["embedding"]:
            try:
                vectors[r["id"]] = json.loads(r["embedding"])
            except Exception:
                pass

    # --- dense leg (only when BOTH a provider is configured AND stored vectors
    # exist in scope — embedding the query when no stored vector can consume it
    # is a wasted API call + latency on every search) --------------------------
    dense: list[str] = []
    if vectors:
        qv = E.embed([query])
        if qv:
            qdim = len(qv[0])
            # Dimension guard: vectors from a different model/provider than the
            # query (mixed corpora after a provider switch) must be skipped —
            # zip() would silently truncate them into garbage similarities.
            scored = [(i, _cosine(qv[0], v)) for i, v in vectors.items()
                      if v and len(v) == qdim]
            scored.sort(key=lambda t: -t[1])
            dense = [i for i, s in scored[: k * 2] if s > 0]

    # --- sparse leg (always available) -------------------------------------
    # BM25's idf can be zero or negative on very small corpora (1-2 docs), so
    # we do not require score > 0. Instead: keep chunks that share at least one
    # meaningful query token, ranked by BM25 score. A query that matches no
    # chunk at all yields no candidates → the agent reports "not found".
    from rank_bm25 import BM25Okapi
    bm = BM25Okapi([_tok(info[i][0]) for i in order])
    sc = bm.get_scores(_tok(query))
    qtok = {t for t in _tok(query) if len(t) >= 3}
    doc_toks = [{t for t in _tok(info[i][0]) if len(t) >= 3} for i in order]
    cand = [i for i in range(len(sc)) if qtok & doc_toks[i]]
    cand.sort(key=lambda i: -sc[i])
    sparse = [order[i] for i in cand[: k * 2]]

    # --- Reciprocal Rank Fusion --------------------------------------------
    rrf: dict[str, float] = {}
    for lst in (dense, sparse):
        for r, i in enumerate(lst):
            rrf[i] = rrf.get(i, 0) + 1 / (60 + r)
    top = sorted(rrf, key=rrf.get, reverse=True)[:k]
    return [{"id": i, "text": info[i][0], "meta": info[i][1], "score": None} for i in top]


def rerank(query, hits, top=6):
    """Cross-encoder rerank (FlashRank). Falls back to fusion order if unavailable."""
    if not hits:
        return []
    try:
        from flashrank import Ranker, RerankRequest
        ranker = Ranker(model_name="ms-marco-MiniLM-L-12-v2")
        res = ranker.rerank(RerankRequest(query=query, passages=[{"id": h["id"], "text": h["text"]} for h in hits]))
        by = {h["id"]: h for h in hits}
        out = []
        for r in res[:top]:
            h = by[r["id"]]; h["score"] = float(r["score"]); out.append(h)
        return out
    except Exception as e:
        print("rerank unavailable:", e)
        return hits[:top]

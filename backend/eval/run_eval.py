"""Offline retrieval + end-to-end evaluation (no network, no API keys, no production).

Runs the REAL pipeline — extraction, chunking, hybrid search (BM25 + dense via
Reciprocal Rank Fusion), the sufficiency gate, the identity-chunk append and
tenant/department authorization — against the synthetic two-organization corpus
in eval/corpus.json and the question set in eval/dataset.jsonl.

Two things are replaced by deterministic local stand-ins (nothing else):

  * LLM       — an EXTRACTIVE stub: it answers from the numbered Sources in
                the prompt by picking the highest question-overlap sentences.
                Answer quality therefore measures retrieval quality, which is
                what this harness exists to measure. The planner returns no
                sub-queries and the rewriter is a no-op, so no LLM crutch can
                hide a retrieval miss.
  * embeddings — a character-trigram hashing embedder (256-dim, L2-normalized)
                plugged in at app.embeddings._openai. The REAL embed() batching,
                the dense leg, the RRF fusion and the dimension guard all run;
                only the HTTP call is replaced. Trigrams give typo/partial-
                overlap robustness similar in kind (not in quality) to a real
                embedding model, so the dense-vs-BM25 delta is measurable.

Metrics per question: Recall@6 and MRR over fused retrieval (retrieval-only,
independent of the LLM), then end-to-end abstention correctness, answer
required-fact accuracy, and unauthorized-leakage checks (evidence, model
prompts and answers are all scanned for the corpus's sentinel strings).

Runtime scenarios beyond the static dataset: document replacement (re-upload
same filename → new facts, old chunks gone), a mid-run new upload, and
multi-file-type ingestion (.docx with tables, .pptx with tables, .csv, .json,
.pdf) verifying the extraction layer end to end.

Usage (from backend/):

    .venv\\Scripts\\python.exe eval\\run_eval.py --mode both
    .venv\\Scripts\\python.exe -m tests.test_retrieval_eval     # CI wrapper w/ thresholds

Exit code is non-zero if any hard invariant fails: protected leakage,
below-threshold questions reaching generation, answerable questions
abstaining, scenario failures. Two things stay in the metric layer (the CI
wrapper pins their thresholds): per-category fact accuracy, and
"incidental overlap" — an unanswerable question whose wording happens to
share >= MIN_EVIDENCE_COVERAGE tokens with an authorized document, so the
calibrated lexical gate passes and generation runs (a known limitation of
any lexical sufficiency gate; leakage is still checked hard there).
"""
import argparse
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# --- env MUST be set before importing the app --------------------------------
_TMP = tempfile.mkdtemp(prefix="omnirag_eval_")
os.environ["OMNIRAG_DATA_DIR"] = _TMP
os.environ["DATABASE_URL"] = ""                     # local sqlite, never production
os.environ["JWT_SECRET"] = "eval-secret-" + "x" * 24
os.environ["ALLOW_ORG_SIGNUP"] = "true"
for _k in ("GROQ_API_KEY", "ANTHROPIC_API_KEY", "GEMINI_API_KEY"):
    os.environ[_k] = ""
# key PRESENCE only — the value is never used because _openai is monkeypatched
os.environ["OPENAI_API_KEY"] = "eval-stub-key"
os.environ["EMBED_PROVIDER"] = "auto"

from app import agent, authz, db as D, embeddings as E, llm, rag  # noqa: E402
from app.config import MIN_EVIDENCE_COVERAGE  # noqa: E402  (same value the gate uses)

EVAL_DIR = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# Deterministic stand-ins
# ---------------------------------------------------------------------------
_TRI = 256


def trigram_embed(texts):
    """Char-trigram hashing embedder: typo/partial-overlap similarity, offline."""
    out = []
    for t in texts:
        v = [0.0] * _TRI
        s = " " + re.sub(r"\s+", " ", t.lower()) + " "
        for i in range(len(s) - 2):
            h = int(hashlib.md5(s[i:i + 3].encode()).hexdigest(), 16)
            v[h % _TRI] += 1.0
        n = sum(x * x for x in v) ** 0.5 or 1.0
        out.append([x / n for x in v])
    return out


E._openai = trigram_embed                       # only the network layer is replaced

CALLS: list[dict] = []                          # every (system, prompt) issued


def _is_answer_call(system: str) -> bool:
    return "fact-checker" not in system and "rewritten search query" not in system \
        and "standalone search queries" not in system


def _extractive_answer(prompt: str) -> str:
    """Deterministic extractive answer from the numbered Sources block.

    Takes the top-5 question-overlap sentences plus the sentence following
    each (policy docs place the answering clause after the matched heading),
    i.e. a slightly generous stand-in for a real grounded LLM. Breadth matters
    for questions whose required facts live in different sentences —
    cross-chunk/multi-doc questions in particular. Ties keep source order
    (stable sort = retrieval rank).
    """
    m = re.search(r"Sources:\n(.*?)\n\nQuestion: (.*)$", prompt, re.S)
    if not m:
        return "I couldn't find this in the documents you have access to."
    src, q = m.group(1), m.group(2)
    qtok = rag._content_tokens(q)
    if not qtok:
        return "I couldn't find this in the documents you have access to."
    # a "source line" is a numbered header, a paragraph/row, or a sentence.
    # Sentence boundaries skip initials ("S. Kumar"), so table rows / names
    # survive intact; paragraphs and rows still split at real sentence ends.
    # Chunk overlap can repeat the SAME sentence across chunks — a real reader
    # wouldn't spend attention twice, so dedupe before ranking (keeps first,
    # i.e. best-ranked occurrence).
    seen, parts = set(), []
    for s in (x.strip() for x in re.split(r"\n|(?<!\b[A-Z])\.\s+|;\s+", src)):
        if len(s) >= 10 and s not in seen:
            seen.add(s)
            parts.append(s)
    order = sorted(range(len(parts)),
                   key=lambda i: -len(qtok & rag._content_tokens(parts[i])))
    keep = []
    for i in order[:5]:
        if not qtok & rag._content_tokens(parts[i]):
            break
        keep.append(i)
        # policy docs put the answering clause AFTER the matched heading
        # ("Non-solicitation. The receiving party agrees … for two years")
        if i + 1 < len(parts):
            keep.append(i + 1)
    top = [parts[i] for i in sorted(set(keep))]
    return ". ".join(top) + ("." if top else "") if top else \
        "I couldn't find this in the documents you have access to."


def _sources_text(prompt: str) -> str:
    """Document-derived text of a model call ONLY.

    The question itself (and its rewriter echo) legitimately contains corpus
    words like tenant names — scanning whole prompts for sentinels would flag
    every cross-tenant question's own wording. Only the ``Sources:`` block is
    document content, so that is what the leakage scan must cover, together
    with the retrieved evidence and the final answer.
    """
    m = re.search(r"Sources:\n(.*?)(?:\n\nQuestion:|\n\nAnswer:|$)", prompt, re.S)
    return m.group(1) if m else ""


def fake_complete(system, prompt, provider=None, temp=0.1):
    CALLS.append({"system": system, "prompt": prompt})
    if "fact-checker" in system:
        return '{"grounded": true, "score": 90, "unsupported": []}', "stub"
    if "rewritten search query" in system:
        return prompt.strip(), "stub"                    # honest no-op rewriter
    if "standalone search queries" in system:
        return "[]", "stub"                              # no planner crutch
    return _extractive_answer(prompt), "stub"


llm.complete = fake_complete

# ---------------------------------------------------------------------------
# Seeding
# ---------------------------------------------------------------------------
TENANTS: dict[str, str] = {}      # "vantor" -> tenant id
USERS: dict[str, dict] = {}       # "van_hr" -> user row dict
DOCS: dict[str, str] = {}         # "hr-handbook.md" -> doc id


def _seed():
    corpus = json.loads((EVAL_DIR / "corpus.json").read_text(encoding="utf-8"))
    D.init()
    with D.db() as c:
        for key in ("vantor", "zentaro"):
            tid = D.uid()
            TENANTS[key] = tid
            c.execute("INSERT INTO tenants(id,name) VALUES(?,?)",
                      (tid, "Vantor Industries" if key == "vantor" else "Zentaro Labs"))
            D.seed_departments(c, tid)
        people = [  # (username, tenant, role, department)
            ("van_admin", "vantor", "admin", "Human Resources"),
            ("van_hr", "vantor", "member", "Human Resources"),
            ("van_eng", "vantor", "member", "Engineering & Tech"),
            ("van_fin", "vantor", "member", "Finance & Operations"),
            ("van_legal", "vantor", "member", "Legal & Compliance"),
            ("van_fin_dir", "vantor", "member", "Finance & Operations"),
            ("zen_hr", "zentaro", "member", "Human Resources"),
            ("zen_admin", "zentaro", "admin", "General"),
        ]
        for uname, tkey, role, dept in people:
            uid = D.uid()
            c.execute("INSERT INTO users(id,username,pw_hash,tenant_id,role,department)"
                      " VALUES(?,?,?,?,?,?)", (uid, uname, "x", TENANTS[tkey], role, dept))
            USERS[uname] = {"id": uid, "username": uname, "tenant_id": TENANTS[tkey],
                            "role": role, "department": dept}
        # organization-visibility docs belong to no department
        for d in corpus["docs"]:
            dept = d.get("department", "")
            if d["visibility"] == "organization":
                dept = ""
            doc_id = D.uid()
            DOCS[d["file"]] = doc_id
            c.execute("INSERT INTO documents(id,tenant_id,filename,visibility,department,"
                      "uploaded_by,chunks) VALUES(?,?,?,?,?,?,?)",
                      (doc_id, TENANTS[d["tenant"]], d["file"], d["visibility"], dept,
                       "eval", 0))
            for acl_user in d.get("acl", []):
                c.execute("INSERT INTO doc_acl(doc_id,user_id) VALUES(?,?)",
                          (doc_id, USERS[acl_user]["id"]))
    for d in corpus["docs"]:
        n = rag.ingest(TENANTS[d["tenant"]], DOCS[d["file"]], d["file"],
                       d["text"].encode("utf-8"))
        with D.db() as c:
            c.execute("UPDATE documents SET chunks=? WHERE id=?", (n, DOCS[d["file"]]))


def _scope(uname: str) -> tuple[dict, list]:
    u = USERS[uname]
    return u, authz.visible_doc_ids(u)


# ---------------------------------------------------------------------------
# Multi-file-type fixtures (extraction layer, incl. table extraction)
# ---------------------------------------------------------------------------
def _pdf_bytes(lines: list[str]) -> bytes:
    def esc(s):
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + \
        " ".join(f"({esc(x)}) Tj T*" for x in lines) + " ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return bytes(out)


def _docx_bytes(paras: list[str], table: list[list[str]]) -> bytes:
    import docx
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    t = d.add_table(rows=len(table), cols=len(table[0]))
    for i, row in enumerate(table):
        for j, cell in enumerate(row):
            t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_bytes(title: str, body: str, table: list[list[str]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(2))
    box.text_frame.text = body
    rows, cols = len(table), len(table[0])
    shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(3.6),
                                   Inches(9), Inches(1.2))
    for i, row in enumerate(table):
        for j, cell in enumerate(row):
            shape.table.cell(i, j).text = cell
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_doc(tkey, filename, data, dept="Engineering & Tech", visibility="department"):
    tid, doc_id = TENANTS[tkey], D.uid()
    with D.db() as c:
        c.execute("INSERT INTO documents(id,tenant_id,filename,visibility,department,"
                  "uploaded_by,chunks) VALUES(?,?,?,?,?,?,?)",
                  (doc_id, tid, filename, visibility, dept, "eval", 0))
    n = rag.ingest(tid, doc_id, filename, data)
    with D.db() as c:
        c.execute("UPDATE documents SET chunks=? WHERE id=?", (n, doc_id))
    DOCS[filename] = doc_id
    return doc_id


# ---------------------------------------------------------------------------
# Metrics
# ---------------------------------------------------------------------------
def _hit_docs(hits):
    return [h["meta"]["doc_id"] for h in hits]


def retrieval_metrics(question, tid, scope, gold_files, k=6):
    """Retrieval-only quality: Recall@k and MRR over fused ranking (+ leakage)."""
    gold = {DOCS[f] for f in gold_files if f in DOCS}
    hits = rag.search(tid, scope, question, k)
    docs = _hit_docs(hits)
    recall = len(gold & set(docs)) / len(gold) if gold else 0.0
    mrr = 0.0
    for rank, did in enumerate(docs, 1):
        if did in gold:
            mrr = 1.0 / rank
            break
    return recall, mrr, hits


def run_question(q, mode_note=""):
    u, scope = _scope(q["user"])
    tid = u["tenant_id"]
    recall, mrr, hits = retrieval_metrics(q["question"], tid, scope, q["gold_docs"])
    CALLS.clear()
    res = agent.run(q["question"], tid, scope, [], None)
    # the agent's own decisive gate coverage (last assess wins across retries)
    cov = None
    for t in res.get("trace", []):
        m = re.search(r"assess: coverage ([0-9.]+)", t)
        if m:
            cov = float(m.group(1))
    ev_text = " ".join(h["text"].lower() for h in hits)
    ans = res["answer"].lower()
    # document-derived prompt content only (never the echoed question wording)
    src_text = " ".join(_sources_text(c["prompt"]).lower() for c in CALLS)
    leak_ev = [p for p in q.get("protected", []) if p.lower() in ev_text]
    leak_pr = [p for p in q.get("protected", [])
               if p.lower() in src_text or p.lower() in ans]
    answered = not res["abstained"]
    facts_ok = True
    if not q["expect_abstain"]:
        for group in q.get("required_facts", []):
            if not any(v.lower() in ans for v in group):
                facts_ok = False
    # Honest abstention classification:
    #   * answerable -> must answer (hard).
    #   * unanswerable with decisive coverage < MIN_EVIDENCE_COVERAGE -> the
    #     calibrated gate MUST block before generation; answering is the exact
    #     production bug class (hard failure).
    #   * unanswerable with coverage >= threshold -> the lexical gate passes
    #     by design (incidental token overlap with an authorized doc). That is
    #     a known gate limitation, measured and reported — the harness must
    #     not pretend a lexical gate can detect it. Leakage stays hard.
    if q["expect_abstain"]:
        gate_blocked = cov is not None and cov < MIN_EVIDENCE_COVERAGE
        abstain_ok = (not answered) if gate_blocked else True
        incidental = gate_blocked is False and answered
    else:
        gate_blocked, incidental = False, False
        abstain_ok = answered                     # answerable questions must answer
    # hard gate contract: a below-threshold question may never reach generation
    gate_ok = (not q["expect_abstain"]) or (not gate_blocked) or \
        not any(_is_answer_call(c["system"]) for c in CALLS)
    return {"id": q["id"], "category": q["category"], "recall": recall, "mrr": mrr,
            "abstain_ok": abstain_ok, "facts_ok": facts_ok, "gate_ok": gate_ok,
            "leak_ev": leak_ev, "leak_pr": leak_pr, "abstained": res["abstained"],
            "coverage": cov, "incidental": incidental,
            "n_hits": len(res["evidence"]), "mode": mode_note}


# ---------------------------------------------------------------------------
# Runtime scenarios (replacement / new upload / file types)
# ---------------------------------------------------------------------------
FORMAT_QUESTIONS = [
    ("docx_table", "van_hr", "What is the desk equipment budget per employee?",
     [["750", "7500"]]),
    ("pptx_table", "van_eng", "Who is the backup database administrator?",
     [["R. Menon"]]),
    ("csv", "van_eng", "How many servers are in the staging environment?",
     [["14"]]),
    ("json", "van_eng", "What is the retry count for the notification worker?",
     [["5", "five"]]),
    ("pdf", "van_hr", "On what date does the parental leave policy review happen?",
     [["march 15", "15 march", "march"]]),
]


def run_scenarios():
    """Replacement, new-upload, and file-type ingestion against the real pipeline."""
    out = []
    _seed_formats()
    for name, user, question, groups in FORMAT_QUESTIONS:
        u, scope = _scope(user)
        CALLS.clear()
        res = agent.run(question, u["tenant_id"], scope, [], None)
        ok = not res["abstained"] and any(
            v.lower() in res["answer"].lower() for grp in groups for v in grp)
        out.append({"scenario": f"format:{name}", "ok": ok,
                    "answer": res["answer"][:120]})

    # --- replacement: same filename, new content ------------------------------
    new_text = ("Travel and Leave Update. Full-time employees now accrue 25 paid "
                "vacation days per year, credited monthly, following a 90-day "
                "probation period. NOTE-REPLACED-9381: this document supersedes "
                "the older handbook excerpt.")
    _replace_doc("vantor", "hr-handbook.md", new_text.encode("utf-8"))
    u, scope = _scope("van_hr")
    CALLS.clear()
    res = agent.run("How many paid vacation days do employees accrue?", u["tenant_id"],
                    scope, [], None)
    new_ok = "25" in res["answer"] and not res["abstained"]
    with D.db() as c:  # the OLD fact must be gone from storage entirely
        stale = c.execute("SELECT COUNT(*) FROM chunks WHERE text LIKE ?",
                          ("%20 paid vacation days per year%",)).fetchone()[0]
    out.append({"scenario": "replacement:new_facts", "ok": new_ok,
                "answer": res["answer"][:120]})
    out.append({"scenario": "replacement:old_chunks_deleted", "ok": stale == 0})

    # --- new upload mid-run ---------------------------------------------------
    _add_doc("vantor", "it-badge.md",
             "IT Badge Policy. A lost access badge must be reported to the IT service "
             "desk within 24 hours. Replacement badges are issued within two working "
             "days at a charge of 500 INR.".encode("utf-8"),
             dept="Engineering & Tech")
    u, scope = _scope("van_eng")
    CALLS.clear()
    res = agent.run("What does the badge policy say about the charge for a lost "
                    "access badge?", u["tenant_id"], scope, [], None)
    ok = not res["abstained"] and "500" in res["answer"]
    out.append({"scenario": "new_upload:immediately_retrievable", "ok": ok,
                "answer": res["answer"][:120]})

    # and it must NOT leak across the tenant boundary (document-derived text only)
    u2, scope2 = _scope("zen_hr")
    CALLS.clear()
    res2 = agent.run("What is the charge for a replacement access badge?",
                     u2["tenant_id"], scope2, [], None)
    foreign = " ".join([res2["answer"].lower()]
                       + [_sources_text(c["prompt"]).lower() for c in CALLS])
    leak = "500" in foreign or "replacement badges" in foreign
    out.append({"scenario": "new_upload:no_cross_tenant_leak", "ok": not leak})
    return out


def _replace_doc(tkey, filename, data):
    """Same flow as main.py upload: old doc + chunks deleted, new id inserted."""
    tid = TENANTS[tkey]
    with D.db() as c:
        for old in c.execute("SELECT id FROM documents WHERE tenant_id=? AND filename=?",
                             (tid, filename)).fetchall():
            rag.delete_doc(tid, old["id"])
            c.execute("DELETE FROM documents WHERE id=?", (old["id"],))
    return _add_doc(tkey, filename, data, dept="Human Resources")


def _seed_formats():
    _add_doc("vantor", "hr-desk-equipment.docx",
             _docx_bytes(["Standard equipment provision rules for office desks."],
                         [["Item", "Budget (INR)"],
                          ["Laptop stand", "250"],
                          ["Ergonomic chair", "12000"],
                          ["Desk equipment per employee", "750"]]),
             dept="Human Resources")
    _add_doc("vantor", "eng-dba-rotation.pptx",
             _pptx_bytes("Database Rotation", "Quarterly DBA rotation announcement.",
                         [["Role", "Primary", "Backup"],
                          ["Database administrator", "S. Kumar", "R. Menon"]]),
             dept="Engineering & Tech")
    _add_doc("vantor", "eng-staging.csv",
             ("environment,servers,notes\n"
              "development,6,shared\n"
              "staging,14,pre-production clones\n"
              "production,40,critical\n").encode("utf-8"),
             dept="Engineering & Tech")
    _add_doc("vantor", "eng-worker-retries.json",
             # spaced keys keep the fact inside an extractable source line while
             # remaining valid JSON; the question asks the same words verbatim
             json.dumps({"service": "notification-worker", "retry count": 5,
                         "backoff seconds": 30}, indent=2).encode("utf-8"),
             dept="Engineering & Tech")
    _add_doc("vantor", "hr-leave-review.pdf",
             _pdf_bytes(["Parental Leave Policy Review Notice.",
                         "The parental leave policy review happens on March 15.",
                         "All managers must submit feedback before that date."]),
             dept="Human Resources")


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def _set_mode(dense: bool):
    E.provider = (lambda: "openai") if dense else (lambda: None)


def run(mode="both", verbose=True) -> dict:
    _seed()
    dataset = [json.loads(l) for l in
               (EVAL_DIR / "dataset.jsonl").read_text(encoding="utf-8").splitlines()
               if l.strip()]
    report = {"modes": {}, "scenarios": [], "failures": [], "fact_misses": []}

    modes = ["bm25", "dense"] if mode == "both" else [mode]
    for label in modes:
        _set_mode(label == "dense")
        rows = [run_question(q, label) for q in dataset]
        cats: dict[str, list] = {}
        for r in rows:
            cats.setdefault(r["category"], []).append(r)
        summary = {}
        for cat, rs in cats.items():
            n = len(rs)
            summary[cat] = {
                "n": n,
                "recall@6": round(sum(r["recall"] for r in rs) / n, 3),
                "mrr": round(sum(r["mrr"] for r in rs) / n, 3),
                "abstain_correct": sum(r["abstain_ok"] for r in rs),
                "facts_correct": sum(r["facts_ok"] for r in rs),
                "incidental_overlap": sum(r["incidental"] for r in rs),
                "leaks": sum(len(r["leak_ev"]) + len(r["leak_pr"]) for r in rs),
                "gate_violations": sum(not r["gate_ok"] for r in rs),
            }
            for r in rs:
                if r["leak_ev"] or r["leak_pr"]:
                    report["failures"].append(
                        f"[{label}] {r['id']}: protected sentinel leaked "
                        f"(evidence={r['leak_ev']} prompts/answers={r['leak_pr']})")
                if not r["abstain_ok"]:
                    report["failures"].append(
                        f"[{label}] {r['id']}: abstention mismatch "
                        f"(got abstained={r['abstained']}, coverage={r['coverage']})")
                if not r["gate_ok"]:
                    report["failures"].append(
                        f"[{label}] {r['id']}: below-threshold question reached "
                        f"generation (coverage={r['coverage']})")
                if not r["facts_ok"]:
                    report["fact_misses"].append(f"[{label}] {r['id']}")
        report["modes"][label] = summary
        if verbose:
            print(f"\n=== Mode: {label} ===")
            print(f"{'category':<22}{'n':>3}{'R@6':>7}{'MRR':>7}"
                  f"{'abstain':>9}{'facts':>7}{'inc':>5}{'leaks':>7}")
            for cat, s in summary.items():
                print(f"{cat:<22}{s['n']:>3}{s['recall@6']:>7}{s['mrr']:>7}"
                      f"{s['abstain_correct']:>4}/{s['n']:<5}"
                      f"{s['facts_correct']:>4}/{s['n']:<3}"
                      f"{s['incidental_overlap']:>5}{s['leaks']:>7}")
            inc = [r for rs in cats.values() for r in rs if r["incidental"]]
            if inc:
                print("  incidental overlap (gate passed on unrelated shared "
                      "tokens; generation ran, leakage still checked):")
                for r in inc:
                    print(f"    {r['id']}  coverage={r['coverage']}")

    _set_mode(True)                                  # scenarios run with dense on
    report["scenarios"] = run_scenarios()
    if verbose:
        print("\n=== Runtime scenarios ===")
        for s in report["scenarios"]:
            print(f"  {'PASS' if s['ok'] else 'FAIL'}  {s['scenario']}"
                  + ("" if s["ok"] else f"  -- {s.get('answer', '')}"))
    for s in report["scenarios"]:
        if not s["ok"]:
            report["failures"].append(f"scenario failed: {s['scenario']}")
    if verbose:
        print(f"\nfailures: {len(report['failures'])}")
        for f in report["failures"]:
            print("  !", f)
        if report["fact_misses"]:
            print(f"fact misses (metric layer, not invariants): "
                  f"{len(report['fact_misses'])}")
            for f in report["fact_misses"]:
                print("  ~", f)
    return report


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--mode", choices=["bm25", "dense", "both"], default="both")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()
    report = run(args.mode, verbose=not args.quiet)
    sys.exit(1 if report["failures"] else 0)


if __name__ == "__main__":
    main()
